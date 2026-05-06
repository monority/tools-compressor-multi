import logging
from pathlib import Path
from typing import Any
import io

logger = logging.getLogger(__name__)


def compress_document(src: Path, dst: Path, opts: dict[str, Any]) -> Path:
    """Compress Office documents by optimizing embedded images and metadata."""
    compress_imgs = opts.get("compress_images", True)
    img_quality = opts.get("image_quality", 50)
    remove_meta = opts.get("remove_metadata", False)
    suffix = src.suffix.lower()
    if suffix == ".docx":
        return _compress_docx(src, dst, compress_imgs, img_quality, remove_meta)
    if suffix == ".xlsx":
        return _compress_xlsx(src, dst, compress_imgs, img_quality, remove_meta)
    raise ValueError(f"Unsupported document format: {suffix}")


def _reencode_image(raw_bytes: bytes, quality: int) -> bytes:
    from PIL import Image

    with Image.open(io.BytesIO(raw_bytes)) as img:
        buf = io.BytesIO()
        fmt = (img.format or "PNG").upper()
        if fmt in ("JPEG", "JPG", "WEBP"):
            if img.mode in ("RGBA", "P"):
                img = img.convert("RGB")
            img.save(buf, format="JPEG" if fmt == "JPG" else fmt, quality=quality, optimize=True)
        else:
            img.save(buf, format=fmt, optimize=True)
        return buf.getvalue()


def _compress_docx(src, dst, compress_imgs, quality, remove_meta):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")

    out = dst.with_suffix(".docx") if dst.suffix != ".docx" else dst
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(src))

    if remove_meta:
        core_props = doc.core_properties
        for prop in ["author", "title", "subject", "keywords", "comments"]:
            try:
                setattr(core_props, prop, "")
            except (AttributeError, TypeError, ValueError) as exc:
                logger.debug("DOCX metadata clear skipped for %s: %s", prop, exc)

    if compress_imgs and quality < 80:
        try:
            for rel in doc.part.rels.values():
                if "image" not in rel.target_ref:
                    continue
                img_part = rel.target_part
                img_bytes = img_part.blob
                try:
                    compressed = _reencode_image(img_bytes, quality)
                    if len(compressed) < len(img_bytes):
                        img_part._blob = compressed
                except (AttributeError, KeyError, OSError, ValueError, TypeError) as exc:
                    logger.debug("DOCX image compression skipped: %s", exc)
        except (AttributeError, KeyError, OSError, ValueError, TypeError) as exc:
            logger.debug("DOCX image walk skipped: %s", exc)

    doc.save(str(out))
    return out


def _compress_xlsx(src, dst, compress_imgs, quality, remove_meta):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    out = dst.with_suffix(".xlsx") if dst.suffix != ".xlsx" else dst
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(str(src))

    if remove_meta:
        wb.properties.creator = ""
        wb.properties.title = ""

    if compress_imgs and quality < 80:
        try:
            for sheet in wb.worksheets:
                for img in getattr(sheet, "_images", []):
                    try:
                        original = img._data()
                        compressed = _reencode_image(original, quality)
                        if len(compressed) < len(original):
                            img._data = lambda data=compressed: data
                    except (AttributeError, KeyError, OSError, ValueError, TypeError) as exc:
                        logger.debug("XLSX image compression skipped: %s", exc)
        except (AttributeError, KeyError, OSError, ValueError, TypeError) as exc:
            logger.debug("XLSX image walk skipped: %s", exc)

    wb.save(str(out))
    return out
